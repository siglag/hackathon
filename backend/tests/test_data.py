import io
import json

from openpyxl import Workbook, load_workbook

from app import artifacts, datasets, db
from app.ingest import canonical, extraction, tabular, validate

db.init_db()


def _canon():
    def s(text, x, y, score=0.99):
        return {"text": text, "bbox": [x, y, x + 60, y + 20], "score": score, "source": "ocr"}
    pages = [{"page": 1, "width": 800, "height": 600, "spans": [
        s("Patient: Jane Testcase MRN: 00999", 10, 10),
        s("Collected: 2026-10-01", 10, 40),
        s("Hemoglobin", 10, 100), s("13.2", 200, 100), s("g/dL", 300, 100), s("13.0-17.0", 400, 100),
        s("Glucose", 10, 140), s("5,8", 200, 140), s("mmol/L", 300, 140), s("3.9-5.5", 400, 140),
        s("Creatinine", 10, 180), s("110", 200, 180, score=0.62), s("umol/L", 300, 180),
    ]}]
    return canonical.build("doc_test", pages, "test")


def test_extraction_keeps_raw_text_and_never_converts():
    ex = extraction.propose(_canon(), "lab_results")
    rows = {r["fields"]["test"]["raw_text"]: r["fields"] for r in ex["rows"]}
    assert rows["Hemoglobin"]["value"]["value"] == 13.2
    assert rows["Hemoglobin"]["value"]["raw_text"] == "13.2"
    # comma decimal: kept as raw text, not silently converted
    assert rows["Glucose"]["value"]["value"] is None and rows["Glucose"]["value"]["raw_text"] == "5,8"
    assert rows["Creatinine"]["reference_range"]["raw_text"] is None  # missing stays unknown
    assert any("Low recognition score" in i for i in rows["Creatinine"]["value"]["issues"])
    assert rows["Hemoglobin"]["value"]["source"]["bbox"] and rows["Hemoglobin"]["value"]["review_state"] == "unreviewed"
    assert ex["header"]["patient_name"]["raw_text"] == "Jane Testcase"
    assert ex["header"]["patient_id"]["raw_text"] == "00999"
    assert rows["Hemoglobin"]["collection_date"]["value"] == "2026-10-01"


def test_model_cannot_introduce_values():
    def fake_model(system, user):
        return {"rows": [
            {"line_id": "p1L2", "test": "Hemoglobin", "value": "14.0", "unit": "g/dL"},  # 14.0 is not on the line
            {"line_id": "p1L3", "test": "Glucose", "value": "5,8", "unit": "mmol/L", "reference_range": "3.9-5.5"},
            {"line_id": "p1L4", "test": "Creatinine", "value": "110", "unit": "umol/L"},
            {"line_id": "nope", "test": "Sodium", "value": "140"},
        ]}
    rows = extraction.model_rows(_canon(), "lab_results", fake_model)
    by = {r["fields"]["test"]["raw_text"]: r for r in rows}
    assert by["Hemoglobin"]["fields"]["value"]["raw_text"] is None  # invented value discarded
    assert "Sodium" not in by


def test_csv_import_preserves_values():
    data = b"test,value,unit\nHbA1c,7.10,%\nLDL,,mmol/L\nNote,=1+1,\nID,007,\n"
    t = tabular.parse_csv(data, "doc1")[0]
    vals = [{k: c["v"] for k, c in r["cells"].items()} for r in t["rows"]]
    assert vals[0]["value"] == 7.1 and t["rows"][0]["cells"]["value"]["r"] == "7.10"
    assert vals[1]["value"] is None  # missing is null, never zero
    assert vals[3]["value"] == "007"  # leading zeros kept as text


def _xlsx_bytes():
    wb = Workbook()
    ws = wb.active
    ws.append(["name", "a", "b", "sum"])
    ws.append(["x", 1, 2, "=B2+C2"])  # no cached value: openpyxl never computes
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_xlsx_formula_without_cache_is_unavailable():
    data = _xlsx_bytes()
    assert validate.sniff(data, "x.xlsx").kind == "xlsx"
    tables, extra = tabular.parse_xlsx(data, "doc2")
    cell = tables[0]["rows"][0]["cells"]["sum"]
    assert cell["v"] is None and "unavailable" in cell


def test_transforms_create_versions_and_artifacts_escape_formulas():
    t = tabular.parse_csv(b"test,value\nB,2\nA,1\nC,=HYPERLINK(\"http://evil\")\n", "doc3")[0]
    did = datasets.create("case_t", "t", t["columns"], t["rows"], ["doc3"], "imported")
    nid = datasets.transform(did, [{"op": "sort", "column": "test"}, {"op": "add_column", "name": "unit", "value": "mmol/L"}])
    d2 = datasets.load(nid)
    assert d2["parent_id"] == did and d2["version"] == 2
    assert [r["cells"]["test"]["v"] for r in d2["rows"]] == ["A", "B", "C"]
    assert len(datasets.load(did)["rows"]) == 3  # input untouched
    aid = artifacts.create("case_t", "xlsx", nid)
    a = db.q1("SELECT * FROM artifacts WHERE id=?", (aid,))
    from app import storage
    wb = load_workbook(io.BytesIO(storage.read(a["storage_ref"])))
    vals = [c.value for c in wb["Data"]["B"]]
    assert "'=HYPERLINK(\"http://evil\")" in vals  # escaped, not a live formula
    assert all(c.data_type != "f" for row in wb["Data"].iter_rows() for c in row)
    csv_id = artifacts.create("case_t", "csv", nid)
    csv_txt = storage.read(db.q1("SELECT storage_ref FROM artifacts WHERE id=?", (csv_id,))["storage_ref"]).decode("utf-8-sig")
    assert "'=HYPERLINK" in csv_txt
    for fmt in ("json", "docx", "pdf", "md"):
        artifacts.create("case_t", fmt, nid)


def test_unsafe_transform_rejected():
    t = tabular.parse_csv(b"a\n1\n", "doc4")[0]
    did = datasets.create("case_t", "t", t["columns"], t["rows"], [], "imported")
    for bad in ([{"op": "exec", "code": "import os"}], [{"op": "filter", "column": "a", "operator": "__import__", "value": 1}],
                [{"op": "sort", "column": "missing"}]):
        try:
            datasets.transform(did, bad)
        except datasets.TransformError:
            continue
        raise AssertionError(f"accepted {bad}")


def test_upload_validation():
    import pytest
    with pytest.raises(validate.Rejected):
        validate.sniff(b"MZ\x90\x00binary", "x.exe")
    with pytest.raises(validate.Rejected):
        validate.sniff(b"{not json", "x.json")
    assert validate.sniff(json.dumps([{"a": 1}]).encode(), "a.json").kind == "json"
