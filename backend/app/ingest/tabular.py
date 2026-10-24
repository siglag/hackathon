"""Direct parsing of XLSX / CSV / JSON (spec 6, 8). No OCR of spreadsheet cells.

Dataset cell format: {"v": typed value, "r": raw text (when it differs),
"src": provenance}. Missing values are ``None``, never zero.

Workbooks are opened twice: with cached values (what Excel last computed) and
with formulas, so a formula lacking a cached result is reported as
*unavailable* rather than recalculated or invented. Macros and external links
are never executed or followed.
"""
from __future__ import annotations

import csv
import datetime as dt
import io
import json
import zipfile
from typing import Any

from ..config import settings


class TabularError(ValueError):
    pass


def _cell(v: Any, raw: str | None = None, src: dict | None = None) -> dict:
    c: dict[str, Any] = {"v": v}
    if raw is not None and raw != (str(v) if v is not None else ""):
        c["r"] = raw
    if src:
        c["src"] = src
    return c


def infer_type(values: list[Any]) -> str:
    kinds = {type(v).__name__ for v in values if v is not None}
    if not kinds:
        return "empty"
    if kinds <= {"int", "float"}:
        return "number"
    if kinds == {"bool"}:
        return "boolean"
    if kinds == {"str"}:
        return "string"
    if kinds <= {"datetime", "date", "str"} and kinds & {"datetime", "date"}:
        return "date" if kinds != {"datetime", "date", "str"} else "mixed"
    return "mixed"


def _jsonable(v: Any) -> Any:
    if isinstance(v, (dt.datetime, dt.date)):
        return v.isoformat()
    if isinstance(v, dt.time):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer() and abs(v) < 2**53:
        return v
    return v


def _unique_names(names: list[Any]) -> list[str]:
    out, seen = [], {}
    for i, n in enumerate(names):
        base = str(n).strip() if n not in (None, "") else f"column_{i + 1}"
        k = base
        while k in seen:
            seen[base] += 1
            k = f"{base}_{seen[base]}"
        seen.setdefault(k, 0)
        out.append(k)
    return out


def _parse_text_value(s: str) -> Any:
    """Type a CSV string conservatively; the raw string is always kept."""
    t = s.strip()
    if t == "":
        return None
    try:
        if t.lstrip("+-").isdigit() and not (len(t.lstrip("+-")) > 1 and t.lstrip("+-").startswith("0")):
            return int(t)
        if t.count(".") == 1 and t.replace(".", "", 1).lstrip("+-").isdigit():
            return float(t)
    except ValueError:
        pass
    return t  # commas, dates, units etc. stay text: never silently reinterpret


def parse_csv(data: bytes, document_id: str) -> list[dict]:
    text = data.decode("utf-8-sig")
    try:
        dialect = csv.Sniffer().sniff(text[:5000], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text), dialect)
    rows = list(reader)
    if not rows:
        raise TabularError("The CSV file has no rows.")
    if len(rows) - 1 > settings.max_sheet_rows:
        raise TabularError(f"The file has more than {settings.max_sheet_rows} rows. Split it into smaller files.")
    header = _unique_names(rows[0])
    out_rows = []
    for ri, r in enumerate(rows[1:], start=2):
        cells = {}
        for ci, name in enumerate(header):
            raw = r[ci] if ci < len(r) else ""
            cells[name] = _cell(_parse_text_value(raw), raw, {"document_id": document_id, "row": ri, "col": ci + 1})
        out_rows.append({"id": f"r{ri}", "cells": cells})
    return [_table("CSV", header, out_rows, notes=[])]


def _table(name: str, header: list[str], rows: list[dict], notes: list[str]) -> dict:
    cols = []
    for h in header:
        cols.append({"name": h, "type": infer_type([r["cells"][h]["v"] for r in rows[:5000]])})
    return {"name": name, "columns": cols, "rows": rows, "notes": notes}


def parse_xlsx(data: bytes, document_id: str) -> tuple[list[dict], dict]:
    from openpyxl import load_workbook

    notes_global: list[str] = []
    wb_vals = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
    wb_form = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
    tables = []
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        names = zf.namelist()
        media = [n for n in names if n.startswith("xl/media/")]
        if any(n.startswith("xl/charts/") for n in names):
            notes_global.append("Charts are not shown in the preview.")
        if any(n.startswith("xl/pivotTables/") for n in names):
            notes_global.append("Pivot tables are not supported; only their cached cell values are shown.")
        if "xl/vbaProject.bin" in names:
            notes_global.append("Macros are ignored and never executed.")
        if any(n.startswith("xl/externalLinks/") for n in names):
            notes_global.append("External links are not followed.")
        notes_global.append("Excel formatting (fonts, colours, conditional formats) is not reproduced.")
    try:
        for ws_v, ws_f in zip(wb_vals.worksheets, wb_form.worksheets):
            if ws_v.sheet_state != "visible":
                notes_global.append(f"Hidden sheet '{ws_v.title}' was imported and is marked hidden.")
            vals = list(ws_v.iter_rows(values_only=True))
            forms = list(ws_f.iter_rows(values_only=True))
            if len(vals) - 1 > settings.max_sheet_rows:
                raise TabularError(f"Sheet '{ws_v.title}' has more than {settings.max_sheet_rows} rows. Split the workbook.")
            # Drop fully empty leading rows
            start = next((i for i, r in enumerate(vals) if any(c not in (None, "") for c in r)), None)
            if start is None:
                tables.append({"name": ws_v.title, "columns": [], "rows": [], "notes": ["Sheet is empty."]})
                continue
            header = _unique_names(vals[start])
            rows = []
            missing_formula = 0
            for ri in range(start + 1, len(vals)):
                rv = vals[ri]
                if not any(c not in (None, "") for c in rv):
                    continue
                rf = forms[ri] if ri < len(forms) else ()
                cells = {}
                for ci, name in enumerate(header):
                    v = rv[ci] if ci < len(rv) else None
                    f = rf[ci] if ci < len(rf) else None
                    src = {"document_id": document_id, "sheet": ws_v.title, "cell": f"{_col_letter(ci + 1)}{ri + 1}"}
                    if isinstance(f, str) and f.startswith("=") and v is None:
                        missing_formula += 1
                        c = {"v": None, "unavailable": "Formula without a cached result; not recalculated.", "formula": f, "src": src}
                    else:
                        c = _cell(_jsonable(v), None, src)
                        if isinstance(f, str) and f.startswith("="):
                            c["formula"] = f
                    cells[name] = c
                rows.append({"id": f"r{ri + 1}", "cells": cells})
            notes = []
            if missing_formula:
                notes.append(f"{missing_formula} formula cell(s) have no cached result and are shown as unavailable.")
            if ws_v.sheet_state != "visible":
                notes.append("Hidden sheet.")
            tables.append(_table(ws_v.title, header, rows, notes))
    finally:
        wb_vals.close()
        wb_form.close()
    return tables, {"notes": notes_global, "embedded_images": media}


def _col_letter(n: int) -> str:
    s = ""
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def parse_json(data: bytes, document_id: str) -> tuple[list[dict], Any]:
    """Preserve the original structure; map compatible arrays of objects to tables."""
    obj = json.loads(data.decode("utf-8-sig"))
    tables = []

    def walk(node: Any, path: str) -> None:
        if len(tables) >= 20:
            return
        if isinstance(node, list) and node and all(isinstance(x, dict) for x in node[:200]):
            keys: list[str] = []
            for x in node[: settings.max_sheet_rows]:
                for k in x:
                    if k not in keys:
                        keys.append(k)
            rows = []
            for i, x in enumerate(node[: settings.max_sheet_rows]):
                cells = {}
                for k in keys:
                    v = x.get(k)
                    if isinstance(v, (dict, list)):
                        cells[k] = {"v": json.dumps(v, ensure_ascii=False), "nested": True, "src": {"document_id": document_id, "path": f"{path}[{i}].{k}"}}
                    else:
                        cells[k] = _cell(v, None, {"document_id": document_id, "path": f"{path}[{i}].{k}"})
                rows.append({"id": f"r{i}", "cells": cells})
            notes = [f"Truncated to {settings.max_sheet_rows} rows."] if len(node) > settings.max_sheet_rows else []
            tables.append(_table(path or "$", [str(k) for k in keys], rows, notes))
            return
        if isinstance(node, dict):
            for k, v in node.items():
                walk(v, f"{path}.{k}" if path else f"$.{k}")
        elif isinstance(node, list):
            for i, v in enumerate(node[:50]):
                walk(v, f"{path}[{i}]")

    walk(obj, "$")
    return tables, obj
